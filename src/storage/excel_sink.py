"""Excel output — 3 sheets per Section 5: Database / Discarded / Audit_Log."""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.storage.models import RawArticle
from src.storage.schema import (
    AUDIT_HEADERS,
    AUDIT_TAB,
    DATABASE_HEADERS,
    DATABASE_TAB,
    DISCARDED_HEADERS,
    DISCARDED_TAB,
)
from src.utils.promo_dates import extract_promo_dates

_HFILL = PatternFill(start_color="FF1F4E78", end_color="FF1F4E78", fill_type="solid")
_HFONT = Font(bold=True, color="FFFFFFFF")


def _iso_week(dt: Optional[datetime]) -> int:
    d = dt or datetime.now(timezone.utc)
    return d.isocalendar().week


def _fmt(dt: Optional[datetime]) -> str:
    if dt is None:
        return ""
    return dt.astimezone(timezone.utc).strftime("%Y-%m-%d")


def _ga(a: RawArticle, attr: str, default=""):
    return getattr(a, attr, default)


def _database_row(a: RawArticle, idx: int) -> list:
    title_raw = a.title_original or ""
    snippet = (a.content_snippet or "").replace("\n", " ").strip()
    summary_rough = snippet[:300]
    start, end = extract_promo_dates(snippet + " " + title_raw)
    cat = ""           # HUMAN
    subcat = ""        # HUMAN
    merged = ""        # auto from cat|subcat (blank Round 1)
    r1 = _ga(a, "_R1", "")
    r2 = _ga(a, "_R2", "")
    sig = _ga(a, "_signal_score", "")
    return [
        idx,                                        # 1 id
        a.url,                                       # 2 url
        title_raw,                                   # 3 hyperlink (rough)
        a.source,                                    # 4 source_name
        _ga(a, "_signal_level", ""),                 # 5 signal_level
        _iso_week(a.published_date or a.crawled_at), # 6 week
        _fmt(a.crawled_at),                          # 7 crawl_date
        title_raw,                                   # 8 title_raw
        _fmt(a.published_date),                      # 9 publish_date
        _ga(a, "_topic_group", "Market Pulse"),      # 10 topic_group
        _ga(a, "_sub_topic_group", "Quốc tế"),       # 11 sub_topic_group
        cat,                                         # 12 category HUMAN
        subcat,                                      # 13 subcategory HUMAN
        merged,                                      # 14 merged_category
        title_raw,                                   # 15 title_normalized rough
        summary_rough,                               # 16 summary rough
        start,                                       # 17 start_date
        end,                                         # 18 end_date
        r1,                                          # 19 R1
        r2,                                          # 20 R2
        sig,                                         # 21 signal_score
        "",                                          # 22 related_vertical HUMAN
    ]


def _style_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    for i, _ in enumerate(headers, 1):
        c = ws.cell(row=1, column=i)
        c.fill = _HFILL
        c.font = _HFONT
        ws.column_dimensions[get_column_letter(i)].width = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def write_excel(
    path: Path,
    kept: Iterable[RawArticle],
    discarded: Iterable[RawArticle] = (),
    audit_rows: Iterable[list] = (),
    start_id: int = 1,
) -> int:
    kept = list(kept)
    discarded = list(discarded)
    kept.sort(key=lambda a: -float(_ga(a, "_signal_score", 0) or 0))

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # Sheet 1 — Database
    ws = wb.active
    ws.title = DATABASE_TAB
    _style_header(ws, DATABASE_HEADERS)
    for i, a in enumerate(kept, start=start_id):
        ws.append(_database_row(a, i))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    # Sheet 2 — Discarded
    wd = wb.create_sheet(DISCARDED_TAB)
    _style_header(wd, DISCARDED_HEADERS)
    for a in discarded:
        wd.append([
            a.url,
            a.source,
            a.title_original or "",
            _fmt(a.published_date),
            _ga(a, "_discard_reason", ""),
            _ga(a, "_matched_keywords", ""),
            bool(_ga(a, "_review_flag", False)),
        ])

    # Sheet 3 — Audit_Log
    wa = wb.create_sheet(AUDIT_TAB)
    _style_header(wa, AUDIT_HEADERS)
    for r in audit_rows:
        wa.append(r)

    wb.save(path)
    return len(kept)
